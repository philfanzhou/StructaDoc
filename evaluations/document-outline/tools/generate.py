"""Generate authored synthetic inputs; expected.json is maintained independently."""
import argparse
from pathlib import Path
from io import BytesIO
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from docx import Document
from docx.shared import Pt

ROOT = Path(__file__).resolve().parents[1]
parser = argparse.ArgumentParser()
parser.add_argument('--output', type=Path, default=ROOT / 'corpus')
CORPUS = parser.parse_args().output
CORPUS.mkdir(parents=True, exist_ok=True)
BODY = [
    'This synthetic paragraph describes ordinary document processing.',
    'Files are retained and their pages are read in the stated order.',
    'The example contains no private information or external source material.',
]

def pdf(name, pages, bookmarks=False, labels=False):
    c = canvas.Canvas(str(CORPUS / f'{name}.pdf'), pagesize=A4, invariant=1)
    c.setTitle(f'StructaDoc synthetic {name}')
    c.setAuthor('StructaDoc contributors')
    for p, headings in enumerate(pages):
        if labels:
            c.addPageLabel(p, style='D', start=p + 1)
        for i, (title, level) in enumerate(headings):
            y = 770 - i * 180
            c.setFont('Helvetica-Bold', 20 if level == 1 else 15)
            c.drawString(60, y, title)
            if bookmarks:
                key = f'p{p}-h{i}'
                c.bookmarkPage(key, fit='FitH', top=y + 24)
                c.addOutlineEntry(title, key, level=level - 1)
            c.setFont('Helvetica', 11)
            for j, line in enumerate(BODY):
                c.drawString(60, y - 30 - j * 17, line)
        c.setFont('Helvetica', 10)
        c.drawCentredString(A4[0] / 2, 35, str(p + 1))
        c.showPage()
    c.save()

pdf('book', [[('Opening', 1), ('Scope', 2)], [('Closing', 1)]], bookmarks=True, labels=True)
pdf('numbered-report', [[('1 Introduction', 1), ('1.1 Constraints', 2)], [('2 Results', 1)]], labels=True)
# A printed table of contents with no PDF bookmarks; labels are explicit PDF metadata.
c = canvas.Canvas(str(CORPUS / 'printed-toc.pdf'), pagesize=A4, invariant=1)
c.setTitle('StructaDoc synthetic printed contents')
c.setAuthor('StructaDoc contributors')
for p, lines in enumerate([
    [('Contents', 20), ('Opening ................................ 2', 12), ('  Scope ................................ 2', 12), ('Closing ................................ 3', 12)],
    [('Opening', 20), ('Scope', 15)],
    [('Closing', 20)],
]):
    c.addPageLabel(p, style='D', start=p + 1)
    for i, (line, size) in enumerate(lines):
        y = 770 - i * (40 if p == 0 else 180)
        c.setFont('Helvetica-Bold' if size > 12 else 'Helvetica', size)
        c.drawString(60, y, line)
        if p > 0:
            c.setFont('Helvetica', 11)
            for j, text in enumerate(BODY):
                c.drawString(60, y - 30 - j * 17, text)
    c.showPage()
c.save()
# No bold/large headings or numbering; these sentences are not outline nodes.
c = canvas.Canvas(str(CORPUS / 'unstructured.pdf'), pagesize=A4, invariant=1)
c.setTitle('StructaDoc synthetic prose')
c.setAuthor('StructaDoc contributors')
c.setFont('Helvetica', 11)
for i, line in enumerate(BODY * 3):
    c.drawString(60, 770 - i * 20, line)
c.save()

doc = Document()
doc.core_properties.title = 'StructaDoc synthetic heading styles'
doc.core_properties.author = 'StructaDoc contributors'
from datetime import datetime, timezone
doc.core_properties.created = datetime(2026, 10, 19, tzinfo=timezone.utc)
doc.core_properties.modified = datetime(2026, 10, 19, tzinfo=timezone.utc)
doc.styles['Normal'].font.size = Pt(11)
for i, (title, level) in enumerate([('Opening', 1), ('Scope', 2), ('Closing', 1)]):
    if i == 2:
        doc.add_page_break()
    doc.add_heading(title, level)
    for text in BODY:
        doc.add_paragraph(text)
stream = BytesIO()
doc.save(stream)
# ZIP timestamps otherwise vary on each generation.
with ZipFile(BytesIO(stream.getvalue())) as src, ZipFile(CORPUS / 'heading-styles.docx', 'w') as dst:
    for name in sorted(src.namelist()):
        info = ZipInfo(name, (2026, 10, 19, 0, 0, 0))
        info.compress_type = ZIP_DEFLATED
        dst.writestr(info, src.read(name))
